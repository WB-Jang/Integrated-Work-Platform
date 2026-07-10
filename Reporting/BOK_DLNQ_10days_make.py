import pandas as pd
import openpyxl


def generate_report(raw, format_file, base_yymmdd, mortgage_loan):
    """BOK 10일 연체 보고서.

    통합 raw(.xlsx, 시트: bf_15/cur_15/cur_17) + 양식(format.xlsx)을 받아
    양식의 '보고서1'/'보고서2' 시트에 최신 데이터를 기입하여 저장한다.
    최종 파일명/경로는 기존 하드코딩 규칙을 그대로 따른다.
    """
    logs = []

    def _log(m):
        logs.append(str(m))

    def _fail(step, e):
        import traceback
        _log(f"[ERROR] '{step}' 단계에서 오류: {e}")
        _log(traceback.format_exc())
        return {'ok': False, 'log': "\n".join(logs), 'outfile': None}

    try:
        _log("[1/5] 입력값 확인 및 통합 raw 로드 중...")
        mortgage_loan = float(mortgage_loan)
        base_yymmdd = str(base_yymmdd)
        output_dir = f"C:/Users/1598505/OneDrive - Standard Chartered Bank/1.보고서(WB)/2.BOK/1.일보_10일보/2.10일보/1.연체율보고서/{base_yymmdd}/(양식2) (10일보 및 월보) 국내은행 연체율 현황({base_yymmdd}기준).xlsx"
        raw_file = pd.read_excel(raw, sheet_name=['bf_15', 'cur_15', 'cur_17'])
        bf_raw_15 = raw_file['bf_15']
        raw_15 = raw_file['cur_15']
        raw_17 = raw_file['cur_17']
    except Exception as e:
        return _fail("통합 raw 로드", e)

    try:
        _log("[2/5] 15번 쿼리 데이터 가공 중...")
        col_list = raw_15.columns.tolist()
        raw_15[col_list[2]] = bf_raw_15[col_list[-2]]

        for col in col_list[2:]:
            raw_15[col] = raw_15[col] / 100000000

        raw_15.iloc[5, -1] = mortgage_loan
        raw_15.iloc[13, -1] = mortgage_loan
    except Exception as e:
        return _fail("15번 쿼리 가공", e)

    try:
        _log("[3/5] 17번 쿼리 데이터 가공 중...")
        col_list_17 = raw_17.columns.tolist()

        for col in col_list_17[2:]:
            raw_17[col] = raw_17[col] / 100000000

        row_no = [1, 6, 8, 13]

        for r in row_no:
            raw_17.iloc[r, 14] = mortgage_loan
        for r in row_no[:2]:
            raw_17.iloc[r, 15] = raw_15.iloc[5, -2]
        for r in row_no[2:]:
            raw_17.iloc[r, 15] = raw_15.iloc[13, -2]

        raw_17_final = raw_17[col_list_17[6:]].drop(index=[0, 6, 7, 13])
        raw_17_final.reset_index(inplace=True)
        raw_17_final.drop(columns='index', inplace=True)
        raw_17_final.iloc[3, 12] = raw_17_final.iloc[0, 12]
        raw_17_final.iloc[8, 12] = raw_17_final.iloc[5, 12]
        raw_17_final.iloc[0, 12] = 0
        raw_17_final.iloc[5, 12] = 0

        raw_17_final.iloc[3, 13] = raw_17_final.iloc[0, 13]
        raw_17_final.iloc[8, 13] = raw_17_final.iloc[5, 13]
        raw_17_final.iloc[0, 13] = 0
        raw_17_final.iloc[5, 13] = 0

        raw_15['기중 정상화 등(E)'] = raw_15['전기준일 연체잔액 (A)'] + raw_15['기간중 신규연체 (B)'] - raw_15['기간중 상각 (C)'] - raw_15['기간중 대환 (D)'] - raw_15['기준일 연체잔액 (F)']
        report2_col_list = raw_15.columns.tolist()
        report2_final_col_list = report2_col_list[:6] + [report2_col_list[-1]] + report2_col_list[6:8]

        raw_15_final = raw_15[report2_final_col_list].drop(index=[0, 8])
        raw_15_final.set_index(['코드', '코드명'], inplace=True)
    except Exception as e:
        return _fail("17번 쿼리 가공", e)

    try:
        _log("[4/5] 양식(format)에 데이터 기입 중...")
        wb = openpyxl.load_workbook(format_file, read_only=False)
        ws1 = wb['보고서1']

        ws1['B4'] = base_yymmdd
        start_row, end_row = 11, 15
        start_col, end_col = 7, 20

        for r_idx, r in enumerate(range(start_row, end_row + 1)):
            for c_idx, c in enumerate(range(start_col, end_col + 1)):
                if r_idx < raw_17_final.iloc[:len(raw_17_final) // 2, :].shape[0] and c_idx < raw_17_final.iloc[:len(raw_17_final) // 2, :].shape[1]:
                    ws1.cell(row=r, column=c, value=raw_17_final.iloc[r_idx, c_idx])

        start_row, end_row = 22, 26
        start_col, end_col = 7, 20
        raw_17_final2 = raw_17_final.iloc[len(raw_17_final) // 2:, :]

        for r_idx, r in enumerate(range(start_row, end_row + 1)):
            for c_idx, c in enumerate(range(start_col, end_col + 1)):
                if r_idx < raw_17_final2.shape[0] and c_idx < raw_17_final2.shape[1]:
                    ws1.cell(row=r, column=c, value=raw_17_final2.iloc[r_idx, c_idx])

        ws2 = wb['보고서2']

        start_row, end_row = 13, 19
        start_col, end_col = 3, 8

        for r_idx, r in enumerate(range(start_row, end_row + 1)):
            for c_idx, c in enumerate(range(start_col, end_col + 1)):
                if r_idx < raw_15_final.iloc[:len(raw_15_final) // 2, :].shape[0] and c_idx < raw_15_final.iloc[:len(raw_15_final) // 2, :].shape[1]:
                    ws2.cell(row=r, column=c, value=raw_15_final.iloc[r_idx, c_idx])

        start_row, end_row = 27, 33
        start_col, end_col = 3, 8
        raw_15_final2 = raw_15_final.iloc[len(raw_15_final) // 2:, :]

        for r_idx, r in enumerate(range(start_row, end_row + 1)):
            for c_idx, c in enumerate(range(start_col, end_col + 1)):
                if r_idx < raw_15_final2.shape[0] and c_idx < raw_15_final2.shape[1]:
                    ws2.cell(row=r, column=c, value=raw_15_final2.iloc[r_idx, c_idx])
    except Exception as e:
        return _fail("양식 데이터 기입", e)

    try:
        _log("[5/5] 결과 저장 중...")
        wb.save(output_dir)
        wb.close()
    except Exception as e:
        return _fail("결과 저장", e)

    _log(f"[완료] BOK 10일 연체 보고서 작성 완료: {output_dir}")
    return {'ok': True, 'log': "\n".join(logs), 'outfile': output_dir}
