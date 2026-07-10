import pandas as pd
import numpy as np
import configs
import openpyxl

from pandas.api.types import CategoricalDtype
from openpyxl import load_workbook

def preprocessing(data):
    data['Tenor Bucket'] = data['Tenor Bucket'].astype(CategoricalDtype(categories=configs.COLUMN_NMS['TENOR_BUCKET_ORDER'], ordered=True))
    data['Segment Classification'] = data['Segment Classification'].astype(CategoricalDtype(categories=configs.COLUMN_NMS['SEGMENT_ORDER'], ordered=True))
    
    for i in range(len(configs.COLUMN_NMS['CRIR_RAW_COLS'])):
        if i == 0:
            data[data.columns[i]]=pd.to_datetime(data.iloc[:,0],format='%Y-%m-%d')
        elif i == 2:
            data['tenor_deriv']=np.where(data[data.columns[i]]<=1,'<= 1 year','> 1 year')
            data['tenor_deriv']=data['tenor_deriv'].astype('string')
        
        elif i == 8:
            data[data.columns[i]]=data[data.columns[i]].astype('string')
            # if 문 내에는 별도의 block이 아니라서, 변수를 선언을 if문 안에서만 하더라도, local 함수 종료시까지는 메모리에 남아있음 -> 차라리 가독성을 위하여 함수 시작하자마자 변수 선언이 유용할 듯
            # 하지만 i를 사용해야해서,,, if문 안에 위치시키자
            # 만약 정말 메모리가 많이 부족하다면 실행 후 즉시 del을 통해서 garbage collector가 작동하도록 코딩할 것
            conditions = [(data[data.columns[i]].isin(configs.CG_GROUPS['Investment grade (CG1-5)'])),(data[data.columns[i]].isin(configs.CG_GROUPS['Sub-investment grade (CG6-8)'])),\
                          (data[data.columns[i]].isin(configs.CG_GROUPS['Sub-investment grade (CG9-11)'])),(data[data.columns[i]].isin(configs.CG_GROUPS['GSAM (CG12)'])),(data[data.columns[i]].isin(configs.CG_GROUPS['GSAM (CG13-14)']))]
            choices = list(configs.CG_GROUPS.keys())
            data['CG_INVESTMENT']=np.select(conditions,choices,default='Unknown')
            data['CG_INVESTMENT']=data['CG_INVESTMENT'].astype('string')
            data['CG_INVESTMENT'] = data['CG_INVESTMENT'].astype(CategoricalDtype(categories=configs.COLUMN_NMS['CG_ORDER'], ordered=True))
        elif i == 12:
             data[data.columns[i]]=np.where(data[data.columns[i]].isin(['Transportation and Storage','Transportation and Storage - Bulker Water Transport',\
                                                                       'Transportation and Storage - Aviation','Transportation and Storage - Container Water Transport',\
                                                                        'Transportation and Storage - Non O&G tanker Water Transport','O&G Tanker - Water Transport']),\
                                                                            'Transportation and Storage',data[data.columns[i]])
        elif i == 15:
            data[data.columns[i]]=data[data.columns[i]]/1000000
        
        else:
            data[data.columns[i]]=data[data.columns[i]].astype('string')
    return data

def overview_pivot(data):
    overview_pivot = pd.pivot_table(
        data
        ,values = 'Net Nominal'
        ,index = ['Scorecard Grouping']
        ,columns = ['Reporting Month']
        ,aggfunc='sum'
        ,fill_value=0)
    
    overview_final = overview_pivot.iloc[0:0].copy()
    overview_final.loc[0] = np.array(overview_pivot.iloc[1,:])+np.array(overview_pivot.iloc[3,:]) 
    overview_final.loc[1] = np.array(overview_pivot.iloc[0,:])+np.array(overview_pivot.iloc[2,:])
    overview_final.loc[2] = np.array(overview_pivot.iloc[4,:])
    
    return overview_final

def overview_seg_pivot(data):
    overview_pivot = pd.pivot_table(
        data
        ,values = 'Net Nominal'
        
        ,columns = ['Reporting Month']
        ,aggfunc='sum'
        ,fill_value=0)
    overview_seg_pivot = pd.pivot_table(
        data[(data['Product Grouping'].isin(['Marketable Securities','Nostros Interbank']))]
        ,values = 'Net Nominal'
        
        ,columns = ['Reporting Month']
        ,aggfunc='sum'
        ,fill_value=0)
    overview_concat = pd.concat([overview_pivot, overview_seg_pivot],axis=0, ignore_index=True)
    
    for i in range(len(overview_concat.columns)):
        overview_concat.loc[2,overview_concat.columns[i]]=overview_concat.iloc[0,i]-overview_concat.iloc[1,i]
    return overview_concat

def ea_pivot(data, ea_purely, ea_non_purely):
    # 기존: input() 으로 CRC 전달값을 직접 입력받음 → 서버/UI 호환을 위해 인자로 주입
    a,b,c,d = ea_purely
    x,y,z,u = ea_non_purely
    ea_pivot = pd.pivot_table(
        data[~(data['CG'].isin(configs.CG_GROUPS['GSAM (CG13-14)']))]
        ,values = 'Net Nominal'
        
        ,columns = ['Reporting Month']
        ,aggfunc='sum'
        ,fill_value=0)
    
    np_res = (np.array([a,b,c,d])/np.array(ea_pivot.iloc[0,:]))*100
    
    ea_pivot.loc['Early Alerts, USDm']=np.array([a,b,c,d])    
    ea_pivot.loc['of which are Non-purely precautionary, USDm']=np.array([x,y,z,u])
    ea_pivot.loc['as a % of credit grade CG 1-11 exposures']=np_res
    ea_pivot.drop(['Net Nominal'],axis=0,inplace=True)
    
    return ea_pivot

def cg_pivot(data):
    print(configs.CG_GROUPS['Investment grade (CG1-5)'][:-2])
    print(configs.CG_GROUPS['Investment grade (CG1-5)'][:-1])
    t1 = pd.pivot_table(data,values = 'Net Nominal',columns = ['Reporting Month'],aggfunc='sum',fill_value=0)
    t2 = pd.pivot_table(data[(\
                              (data['Scorecard Grouping'].isin(['Corporate','NBFI','OTH']))&(data['CG'].isin(configs.CG_GROUPS['Investment grade (CG1-5)']))\
                              |(data['Scorecard Grouping'].isin(['Bank']))&(data['CG'].isin(configs.CG_GROUPS['Investment grade (CG1-5)'][:-2]))\
                              |(data['Scorecard Grouping'].isin(['Sovereign']))&(data['CG'].isin(configs.CG_GROUPS['Investment grade (CG1-5)'][:-1]))\
                                )],values = 'Net Nominal',columns = ['Reporting Month'],aggfunc='sum',fill_value=0)
    t3 = pd.pivot_table(data[(\
                              (data['Scorecard Grouping'].isin(['Corporate','NBFI','OTH']))&(data['CG'].isin(configs.CG_GROUPS['Sub-investment grade (CG6-8)']+configs.CG_GROUPS['Sub-investment grade (CG9-11)']))\
                              |(data['Scorecard Grouping'].isin(['Bank']))&(data['CG'].isin(configs.CG_GROUPS['Sub-investment grade (CG6-8)']+configs.CG_GROUPS['Sub-investment grade (CG9-11)']+['5A','5B']))\
                              |(data['Scorecard Grouping'].isin(['Sovereign']))&(data['CG'].isin(configs.CG_GROUPS['Sub-investment grade (CG6-8)']+configs.CG_GROUPS['Sub-investment grade (CG9-11)']+['5B']))\
                                )],values = 'Net Nominal',columns = ['Reporting Month'],aggfunc='sum',fill_value=0)
    t4 = pd.pivot_table(data[\
                              (data['CG'].isin(configs.CG_GROUPS['GSAM (CG12)']))\
                              ],values = 'Net Nominal',columns = ['Reporting Month'],aggfunc='sum',fill_value=0)
    t5 = pd.pivot_table(data[\
                              (data['CG'].isin(configs.CG_GROUPS['GSAM (CG13-14)']))\
                              ],values = 'Net Nominal',columns = ['Reporting Month'],aggfunc='sum',fill_value=0)
    cg_table = pd.concat([t1,t2,t3,t4,t5],axis=0, ignore_index=True)
    
    cg_table = cg_table[[cg_table.columns[0],cg_table.columns[1],cg_table.columns[-1]]]
    cg_table.loc['Not Graded'] = np.array([0,0,0])
    return cg_table

def tenor_pivot(data):
    data = data[data['Reporting Month'].isin([np.sort(data['Reporting Month'].unique())[0],np.sort(data['Reporting Month'].unique())[1],np.sort(data['Reporting Month'].unique())[-1]])]
    data['Segment Classification'] = data['Segment Classification'].astype(CategoricalDtype(categories=configs.COLUMN_NMS['SEGMENT_ORDER'], ordered=True))
    t = pd.pivot_table(data,values = 'Net Nominal',index='Segment Classification',columns = ['Reporting Month'],aggfunc='sum',fill_value=0)

    t1 = pd.pivot_table(data[data['Tenor Bucket']=='<=1 year'],values = 'Net Nominal',index='Segment Classification',columns = ['Tenor Bucket','Reporting Month'],aggfunc='sum',fill_value=0)
    t2 = pd.pivot_table(data[data['Tenor Bucket']=='1-3 years'],values = 'Net Nominal',index='Segment Classification',columns = ['Tenor Bucket','Reporting Month'],aggfunc='sum',fill_value=0)
    t3 = pd.pivot_table(data[data['Tenor Bucket']=='3-5 years'],values = 'Net Nominal',index='Segment Classification',columns = ['Tenor Bucket','Reporting Month'],aggfunc='sum',fill_value=0)
    t4 = pd.pivot_table(data[data['Tenor Bucket']=='>5 years'],values = 'Net Nominal',index='Segment Classification',columns = ['Tenor Bucket','Reporting Month'],aggfunc='sum',fill_value=0)

    tenor_t = pd.concat([t1/t,t2/t,t3/t,t4/t],axis=1,ignore_index=True)
    tenor_t = tenor_t*100

    t = pd.pivot_table(data,values = 'Net Nominal',columns = ['Reporting Month'],aggfunc='sum',fill_value=0)
    t1 = pd.pivot_table(data[data['Tenor Bucket']=='<=1 year'],values = 'Net Nominal',columns = ['Tenor Bucket','Reporting Month'],aggfunc='sum',fill_value=0)
    t2 = pd.pivot_table(data[data['Tenor Bucket']=='1-3 years'],values = 'Net Nominal',columns = ['Tenor Bucket','Reporting Month'],aggfunc='sum',fill_value=0)
    t3 = pd.pivot_table(data[data['Tenor Bucket']=='3-5 years'],values = 'Net Nominal',columns = ['Tenor Bucket','Reporting Month'],aggfunc='sum',fill_value=0)
    t4 = pd.pivot_table(data[data['Tenor Bucket']=='>5 years'],values = 'Net Nominal',columns = ['Tenor Bucket','Reporting Month'],aggfunc='sum',fill_value=0)

    tenor_t2 = pd.concat([t1/t,t2/t,t3/t,t4/t],axis=1,ignore_index=True)
    tenor_f = pd.concat([tenor_t,tenor_t2],axis=0)
    tenor_f = tenor_f.rename(index={'Net Nominal':'Total'})

    avg = pd.pivot_table(data,values='Residual Maturity (Years)',index='Segment Classification', columns=['Reporting Month'],aggfunc='mean',fill_value=0)
    avg2 = pd.pivot_table(data,values='Residual Maturity (Years)', columns=['Reporting Month'],aggfunc='mean',fill_value=0)
    avg_f=pd.concat([avg,avg2],axis=0)
    avg_f = avg_f.rename(index={'Residual Maturity (Years)':'Total'})

    tenor_avg = pd.concat([tenor_f,avg_f],axis=1,ignore_index=True)

    return tenor_avg

def industry_pivot(data):
    data = data[data['Reporting Month'].isin([np.sort(data['Reporting Month'].unique())[0],np.sort(data['Reporting Month'].unique())[1],np.sort(data['Reporting Month'].unique())[-1]])]
    industry_t = pd.pivot_table(data[data['Scorecard Grouping'].isin(['Corporate','OTH'])],values='Net Nominal',index='Industry Limits Classification',columns='Reporting Month',aggfunc='sum',fill_value=0)
    industry_t = industry_t.sort_values(by=np.sort(data['Reporting Month'].unique())[-1],ascending=False)

    top10 = industry_t.iloc[:10]
    others = industry_t.iloc[10:]

    others_sum = others.sum(numeric_only=True)
    new_row = pd.Series(index=industry_t.columns,dtype='float')
    new_row.update(others_sum)

    industry_f = pd.concat([top10,pd.DataFrame([new_row])],axis=0)
    industry_f = industry_f.rename(index={0:'Others'})

    total_value = industry_f.sum(numeric_only=True)
    industry_f.loc['Total'] = total_value

    industry_f['QoQ% Change'] = (industry_f.iloc[:,-1] - industry_f.iloc[:,1])/industry_f.iloc[:,1]


    data2 = data[data['Reporting Month']==np.sort(data['Reporting Month'].unique())[-1]]
    

    industry_cg = pd.pivot_table(\
                            data2[(data2['Scorecard Grouping'].isin(['Corporate','OTH']))&(data2['CG_INVESTMENT'].isin(['Investment grade (CG1-5)','GSAM (CG12)','GSAM (CG13-14)']))],\
                                values='Net Nominal',index='Industry Limits Classification',columns=['CG_INVESTMENT'],aggfunc='sum',fill_value=0)
    industry_cg = industry_cg.rename(columns={'Sub-investment grade (CG6-8)':'Sub-investment grade (<= 1 year)','Sub-investment grade (CG9-11)':'Sub-investment grade (> 1 year)'})

    industry_cg2 = pd.pivot_table(\
                            data2[(data2['Scorecard Grouping'].isin(['Corporate','OTH']))&(data2['CG_INVESTMENT'].isin(['Sub-investment grade (CG6-8)','Sub-investment grade (CG9-11)']))],\
                                values='Net Nominal',index='Industry Limits Classification',columns=['tenor_deriv'],aggfunc='sum',fill_value=0)
    
    industry_cg['Sub-investment grade (<= 1 year)'].update(industry_cg2['<= 1 year'])
    industry_cg['Sub-investment grade (> 1 year)'].update(industry_cg2['> 1 year'])
    
    
    industry_cg1 = industry_cg[industry_cg.index.isin(industry_f.index)]
    industry_cg2 = industry_cg[~(industry_cg.index.isin(industry_f.index))]
    others_sum = industry_cg2.sum(numeric_only=True)
    
    industry_cg3 = pd.Series(index=industry_cg2.columns,dtype='float')
    industry_cg3.update(others_sum)
    industry_cg3 = pd.DataFrame([industry_cg3])
    industry_cg3 = industry_cg3.rename(index={0:'Others'})
    
    industry_cg_f = pd.concat([industry_cg1,industry_cg3],axis=0)
    total_value = industry_cg_f.sum(numeric_only=True)
    industry_cg_f.loc['Total'] = total_value

    industry_ratio = (industry_cg_f.div(industry_f[industry_f.columns[-2]], axis=0))
    
    industry_ff = pd.concat([industry_f,industry_ratio],axis=1,ignore_index=False)
              

    return industry_ff

def collateral(data):
    data = data[configs.COLUMN_NMS['CLLTRL_COLS']]
    
    for i in range(len(configs.COLUMN_NMS['CLLTRL_COLS'])):
        if i == 0:
            data[data.columns[i]]=data[data.columns[i]].astype('string')
            conditions = [(data[data.columns[i]].isin(configs.CG_GROUPS_CLLTRL['Investment grade (CG1-5)'])),\
                          (data[data.columns[i]].isin(configs.CG_GROUPS_CLLTRL['Sub-investment grade (CG till 11C)'])),\
                          (data[data.columns[i]].isin(configs.CG_GROUPS_CLLTRL['GSAM (CG12)'])),\
                          (data[data.columns[i]].isin(configs.CG_GROUPS_CLLTRL['GSAM (CG13-14)']))]
            choices = list(configs.CG_GROUPS_CLLTRL.keys())
            data['CG_INVESTMENT']=np.select(conditions,choices,default='Unknown')
            data['CG_INVESTMENT']=data['CG_INVESTMENT'].astype('string')
            data['CG_INVESTMENT'] = data['CG_INVESTMENT'].astype(CategoricalDtype(categories=configs.COLUMN_NMS['CG_ORDER_CLLTRL'], ordered=True))

        else:
            data[data.columns[i]]=data[data.columns[i]]/1000000
    collateral = pd.pivot_table(data,index='CG_INVESTMENT',values=['Ccmv','Net Nominal'],aggfunc='sum',fill_value=0)
    total_value = collateral.sum(numeric_only=True)
    collateral.loc['Total']=total_value
    collateral['ratio'] = (collateral['Ccmv']/collateral['Net Nominal'])
    collateral_f = collateral.iloc[:,-1]
    

    clltrl_sum = data[configs.COLUMN_NMS['CLLTRL_COLS'][3:]].sum(numeric_only=True)
    clltrl_row = pd.Series(index=configs.COLUMN_NMS['CLLTRL_COLS'][3:],dtype='float')
    clltrl_row.update(clltrl_sum)
    total_value = clltrl_row.sum(numeric_only=True)
    clltrl_row['Total'] = total_value
    
    clltrl_row['Guarantee'] = clltrl_row['Guarantee'] + clltrl_row['Securities']
    clltrl_row['Others'] = clltrl_row['Commodities & SIP Repo'] + clltrl_row['Other Collateral'] + clltrl_row['Plant, Machinery & Other stock in WC']
    clltrl_row.drop(['Securities','Commodities & SIP Repo','Other Collateral','Plant, Machinery & Other stock in WC'],inplace=True)
    
    clltrl_row = clltrl_row.reindex(configs.COLUMN_NMS['CLLTRL_ORDER'])
    

    return clltrl_row,collateral_f


def policy_exception(data, base_ym):
    # 기존: base_ym 을 input() 으로 입력 → 인자로 주입
    date_index = pd.date_range(end=pd.to_datetime(base_ym), periods=12, freq='MS')
    date_index = date_index.strftime('%Y%m').tolist()
    
    data.drop_duplicates(subset=['LEID'],keep='first',inplace=True) # keep='last'
    data['Approved Date'] = pd.to_datetime(data['Approved Date'], dayfirst=True)
    data['Approved_mnth'] = data['Approved Date'].dt.strftime('%Y%m')

    pivot_significant = pd.pivot_table(data[(data['BCA Type']=='New')&(data['Approved_mnth']==date_index[-1])&(data["Client Total CAT1&CAT2('000)"]>=1000000)],\
                                       index=['LEID','Customer Name'],values="Client Total CAT1&CAT2('000)",aggfunc='sum',fill_value=0)
    

    data['Exception Code']=data['Exception Code'].str.split(',')
    data = data.explode('Exception Code')
    data['Exception Code']=data['Exception Code'].str.strip()

    pivot_exception = pd.pivot_table(data[(data['Approved_mnth'].isin(date_index))&(~data['Exception Code'].isnull())],\
                                       columns=['Exception Code'],aggfunc='size',fill_value=0)
    

    data['Approval Level Underwriting']=data['Approval Level Underwriting'].str.split(r',|;|-',regex=True)
    data = data.explode('Approval Level Underwriting')
    data['Approval Level Underwriting']=data['Approval Level Underwriting'].str.strip()

    pivot_CPG = pd.pivot_table(data[(data['Approved_mnth'].isin(date_index))&(~data['Approval Level Underwriting'].isnull())&(data['Compliant Country Underwriting']=='No')],\
                                       columns=['Approval Level Underwriting'],aggfunc='size',fill_value=0)
    

    # print(data['Approved_mnth'])
    return pivot_significant,pivot_exception,pivot_CPG
def _annex_corp(data_corp):
    conditions = [(data_corp['Final WACG'].isin(configs.ANNEX3_CG_GROUPS['CG 1-3'])),(data_corp['Final WACG'].isin(configs.ANNEX3_CG_GROUPS['CG 4'])),\
                          (data_corp['Final WACG'].isin(configs.ANNEX3_CG_GROUPS['CG 5'])),(data_corp['Final WACG'].isin(configs.ANNEX3_CG_GROUPS['CG 6'])),\
                            (data_corp['Final WACG'].isin(configs.ANNEX3_CG_GROUPS['CG 7-8'])),(data_corp['Final WACG'].isin(configs.ANNEX3_CG_GROUPS['CG 9-10'])),\
                            (data_corp['Final WACG'].isin(configs.ANNEX3_CG_GROUPS['CG 11'])),(data_corp['Final WACG'].isin(configs.ANNEX3_CG_GROUPS['CG 12']))]
    choices = list(configs.ANNEX3_CG_GROUPS.keys())
    data_corp['WACG BAND']=np.select(conditions,choices,default='Unknown')
    data_corp['WACG BAND']=data_corp['WACG BAND'].astype('string')
    
    pivot_tmp = pd.pivot_table(data_corp, index=['Client Group','WACG BAND'],values='Final exposure',aggfunc='sum',fill_value=0)
    
    ratio_corp = pd.pivot_table(pivot_tmp, index=['WACG BAND'], aggfunc='size')
    ratio_corp = ratio_corp/len(pivot_tmp)*20
    ratio_corp = ratio_corp.round(0).astype('int')
    if ratio_corp.sum() != 20:
        ratio_corp[0]+=(20 - ratio_corp.sum())
    ratio_corp = ratio_corp.reindex(['CG 1-3','CG 4', 'CG 5', 'CG 6', 'CG 7-8', 'CG 9-10','CG 11'])
    ratio_corp = ratio_corp.fillna(0) 
    

    data_corp = data_corp[['Client Group','WACG BAND','Final WACG','Final exposure']]
    data_corp = data_corp.groupby(['Client Group','WACG BAND','Final WACG'], as_index=False)['Final exposure'].sum()
    
    df_list = []
    for i in range(len(ratio_corp)):
        df_list.append(data_corp[data_corp['WACG BAND']==ratio_corp.keys()[i]].sort_values(by='Final exposure', ascending=False)[:ratio_corp[i]])
    
    corp_list_f = pd.concat(df_list, ignore_index=True)
    return corp_list_f

def _annex_fi(data_fi):
    conditions = [(data_fi['Final WACG'].isin(configs.ANNEX3_CG_FI_GROUPS['CG 1-3'])),(data_fi['Final WACG'].isin(configs.ANNEX3_CG_FI_GROUPS['CG 4-5A'])),\
                          (data_fi['Final WACG'].isin(configs.ANNEX3_CG_FI_GROUPS['CG 5B-6'])),\
                            (data_fi['Final WACG'].isin(configs.ANNEX3_CG_FI_GROUPS['CG 7-8'])),(data_fi['Final WACG'].isin(configs.ANNEX3_CG_FI_GROUPS['CG 9-10'])),\
                            (data_fi['Final WACG'].isin(configs.ANNEX3_CG_FI_GROUPS['CG 11'])),(data_fi['Final WACG'].isin(configs.ANNEX3_CG_FI_GROUPS['CG 12']))]
    choices = list(configs.ANNEX3_CG_FI_GROUPS.keys())
    data_fi['WACG BAND']=np.select(conditions,choices,default='Unknown')
    data_fi['WACG BAND']=data_fi['WACG BAND'].astype('string')
    
    pivot_tmp = pd.pivot_table(data_fi, index=['Client Group','WACG BAND'],values='Final exposure',aggfunc='sum',fill_value=0)
    
    ratio_fi = pd.pivot_table(pivot_tmp, index=['WACG BAND'], aggfunc='size')
    ratio_fi = ratio_fi/len(pivot_tmp)*20
    ratio_fi = ratio_fi.round(0).astype('int')
    
    if ratio_fi.sum() != 20:
        ratio_fi[0]+=(20 - ratio_fi.sum())
    ratio_fi = ratio_fi.reindex(['CG 1-3','CG 4-5A', 'CG 5-6B', 'CG 7-8', 'CG 9-10','CG 11'])
    ratio_fi = ratio_fi.fillna(0) 
    ratio_fi = ratio_fi.round(0).astype('int')
    
    data_fi = data_fi[['Client Group','WACG BAND','Final WACG','Final exposure']]
    data_fi = data_fi.groupby(['Client Group','WACG BAND','Final WACG'], as_index=False)['Final exposure'].sum()
    
    df_list = []
    for i in range(len(ratio_fi)):
        print(ratio_fi[i], type(ratio_fi[i]))
        df_list.append(data_fi[data_fi['WACG BAND']==ratio_fi.keys()[i]].sort_values(by='Final exposure', ascending=False)[:ratio_fi[i]])
    
    fi_list_f = pd.concat(df_list, ignore_index=True)
    return fi_list_f
def annex(data):
    
    data_corp = data[(data['Group Segment2'].isin(['Corporate','Others']))&(~data['Customer Sub Segment'].isin(['Public Pension Funds','Real Money Funds']))\
                     &(~data['Client Group'].isin(['SOUTH KOREA GOVERNMENT']))]
    data_corp['Final exposure'] = data_corp['Final exposure']/1000000
    data_fi = data[data['Group Segment2'].isin(['FI'])]
    data_fi['Final exposure'] = data_fi['Final exposure']/1000000
    corp_f = _annex_corp(data_corp)
    fi_f = _annex_fi(data_fi)
    return corp_f, fi_f


# ──────────────────────────────────────────────────────────────────────────────
# 보고서 생성 진입점 (form_fill 패턴)
#
# 기존 스크립트 본문(하드코딩 경로 Z:/..., input() 호출)을 함수로 리팩토링했다.
# raw 엑셀(다중 시트)에서 집계 → 폼 템플릿(.xlsx, 워크시트 'new')의 셀에 직접 기입
# → 결과 파일로 저장한다. 각 섹션은 방어적으로 처리되어, 특정 시트/설정이 없어도
# 가능한 섹션만 채우고 나머지는 건너뛴다(로그로 사유 표시).
# ──────────────────────────────────────────────────────────────────────────────
def generate_crir(raw_path, form_path, output_path,
                  dates, ea_purely, ea_non_purely, base_ym, log_fn=print):
    """CRIR 보고서 생성.

    raw_path     : 원시 데이터 .xlsx (시트: Sheet2[, Collateral, BCA CIB/FI/CC, DATA])
    form_path    : 결과를 채울 폼 템플릿 .xlsx (워크시트 'new')
    output_path  : 채워진 결과 .xlsx 저장 경로
    dates        : 기간 라벨 4개 리스트 (예: ['Apr25','Jan26','Mar26','Apr26'])
    ea_purely / ea_non_purely : Early Alerts USDm 4개 정수 리스트
    base_ym      : 기준년월 'YYYY-MM' (policy_exception 용)
    log_fn       : 진행 로그 콜백 (기본 print)
    """
    raw = pd.read_excel(raw_path, sheet_name='Sheet2',
                        usecols=configs.COLUMN_NMS['CRIR_RAW_COLS'])
    raw = preprocessing(raw)
    log_fn(f"[INFO] raw_data(Sheet2) 로드/전처리 완료: {len(raw)}행")

    wb = openpyxl.load_workbook(form_path, read_only=False)
    if 'new' not in wb.sheetnames:
        raise ValueError(f"폼 템플릿에 'new' 워크시트가 없습니다. (시트: {wb.sheetnames})")
    ws = wb['new']

    # ── 기간 라벨 헤더 기입 ────────────────────────────────────────────
    row_col_list = [[3, 4, 7], [19, 4, 7], [44, 4, 9], [55, 3, 5], [55, 6, 8],
                    [55, 9, 11], [55, 12, 14], [55, 15, 17], [120, 3, 5]]

    def date_change(target_row, start_col_merge, end_col):
        # target_row>32 인 행은 3개 라벨(0,1,3), 그 외는 4개 라벨.
        # 44행은 병합 셀이라 2칸 간격, 그 외는 1칸 간격.
        _dates = [dates[0], dates[1], dates[3]] if target_row > 32 else list(dates)
        for i, value in enumerate(_dates):
            current_col = start_col_merge + (2 * i if target_row == 44 else i)
            if current_col <= end_col:
                ws.cell(row=target_row, column=current_col, value=value)

    for r, c, e in row_col_list:
        date_change(r, c, e)

    def _fill(df, start_row, end_row, start_col, end_col, col_step=1, is_series=False):
        """DataFrame/Series 값을 지정 셀 범위에 기입."""
        for r_idx, r in enumerate(range(start_row, end_row + 1)):
            if is_series:
                if r_idx < df.shape[0]:
                    ws.cell(row=r, column=start_col, value=df.iloc[r_idx])
            else:
                for c_idx, c in enumerate(range(start_col, end_col + 1, col_step)):
                    if r_idx < df.shape[0] and c_idx < df.shape[1]:
                        ws.cell(row=r, column=c, value=df.iloc[r_idx, c_idx])

    # ── 핵심 섹션 (Sheet2 기반) — 각 섹션 실패해도 다음 섹션 진행 ────────
    try:
        _fill(overview_pivot(raw), 7, 9, 4, 7)
        log_fn("[INFO] overview 기입 완료")
    except Exception as e:
        log_fn(f"[주의] overview 실패: {e}")
    try:
        _fill(overview_seg_pivot(raw), 21, 23, 4, 7)
        log_fn("[INFO] overview_seg 기입 완료")
    except Exception as e:
        log_fn(f"[주의] overview_seg 실패: {e}")
    try:
        _fill(ea_pivot(raw, ea_purely, ea_non_purely), 24, 26, 4, 7)
        log_fn("[INFO] Early Alerts 기입 완료")
    except Exception as e:
        log_fn(f"[주의] Early Alerts 실패: {e}")
    try:
        _fill(cg_pivot(raw), 46, 51, 4, 9, col_step=2)
        log_fn("[INFO] CG 기입 완료")
    except Exception as e:
        log_fn(f"[주의] CG 실패: {e}")
    try:
        _fill(tenor_pivot(raw), 56, 61, 3, 17)
        log_fn("[INFO] tenor 기입 완료")
    except Exception as e:
        log_fn(f"[주의] tenor 실패: {e}")
    try:
        _fill(industry_pivot(raw).reset_index(), 122, 133, 2, 11)
        log_fn("[INFO] industry 기입 완료")
    except Exception as e:
        log_fn(f"[주의] industry 실패: {e}")

    # ── Collateral 시트 (선택) ──────────────────────────────────────────
    try:
        raw_collateral = pd.read_excel(raw_path, sheet_name='Collateral')
        clltrl_row, clltrl = collateral(raw_collateral)
        _fill(clltrl_row, 137, 142, 4, 4, is_series=True)
        _fill(clltrl, 143, 147, 4, 4, is_series=True)
        log_fn("[INFO] collateral 기입 완료")
    except Exception as e:
        log_fn(f"[주의] collateral 건너뜀: {e}")

    # ── 정책 예외 (BCA 시트들, 선택) — 폼 좌표 미지정이라 계산/검증만 수행 ──
    for sheet in ('BCA CIB', 'BCA FI', 'BCA CC'):
        try:
            df_bca = pd.read_excel(raw_path, sheet_name=sheet)
            policy_exception(df_bca, base_ym)
            log_fn(f"[INFO] policy_exception({sheet}) 계산 완료")
        except Exception as e:
            log_fn(f"[주의] policy_exception({sheet}) 건너뜀: {e}")

    # ── Annex (DATA 시트, 선택) — configs.ANNEX3_* 설정 필요 ────────────
    try:
        annex3 = pd.read_excel(
            raw_path, sheet_name='DATA',
            usecols=['Client Group', 'Group Segment2', 'Customer Sub Segment',
                     'Final WACG', 'Final exposure'])
        corp_f, fi_f = annex(annex3)
        _fill(corp_f, 193, 212, 3, 6)
        _fill(fi_f, 215, 234, 3, 6)
        log_fn("[INFO] annex 기입 완료")
    except Exception as e:
        log_fn(f"[주의] annex 건너뜀: {e}")

    wb.save(output_path)
    wb.close()
    log_fn(f"[INFO] 폼 템플릿 저장 완료: {output_path}")
    return output_path


def generate_report(raw, format_file, base_ym, dates, ea_purely, ea_non_purely):
    """CRIR 보고서 — 표준 generate_report 진입점.

    통합 raw(.xlsx) + 양식(format.xlsx, 'new' 시트) + 기준년월(base_ym, YYYY-MM)
    + 기간 라벨 4개(dates) + Early Alerts 값(ea_purely/ea_non_purely)을 받아
    양식을 채워 저장한다. 각 섹션 오류는 generate_crir 내부에서 단계별로 로그에
    남으며, 최종 파일명/경로는 기존 하드코딩 규칙을 그대로 따른다.
    """
    logs = []

    def _log(m):
        logs.append(str(m))

    try:
        base_ym = str(base_ym)
        yyyymm = base_ym.replace('-', '').replace('.', '').replace('/', '')
        outfile = f"C:/Users/1598505/OneDrive - Standard Chartered Bank/1.보고서(WB)/3.Internal/8.CRIR/{yyyymm}/overview_{yyyymm[2:]}.xlsx"
        _log(f"[시작] CRIR 보고서 생성 (기준: {base_ym})")
        generate_crir(
            raw_path=raw, form_path=format_file, output_path=outfile,
            dates=dates, ea_purely=ea_purely, ea_non_purely=ea_non_purely,
            base_ym=base_ym, log_fn=_log,
        )
    except Exception as e:
        import traceback
        _log(f"[ERROR] 'CRIR 보고서 생성' 단계에서 오류: {e}")
        _log(traceback.format_exc())
        return {'ok': False, 'log': "\n".join(logs), 'outfile': None}

    _log(f"[완료] CRIR 보고서 작성 완료: {outfile}")
    return {'ok': True, 'log': "\n".join(logs), 'outfile': outfile}


if __name__ == '__main__':
    # 로컬 단독 실행용 예시 (실제 경로/파라미터로 교체해 사용)
    import sys
    if len(sys.argv) >= 3:
        generate_crir(
            raw_path=sys.argv[1],
            form_path=sys.argv[2],
            output_path=sys.argv[3] if len(sys.argv) > 3 else 'CRIR_filled.xlsx',
            dates=['Apr25', 'Jan26', 'Mar26', 'Apr26'],
            ea_purely=[0, 0, 0, 0],
            ea_non_purely=[0, 0, 0, 0],
            base_ym='2026-04',
        )
    else:
        print("usage: python CRIR.py <raw.xlsx> <form.xlsx> [output.xlsx]")
