import pandas as pd
import openpyxl
import configs
import numpy as np

from datetime import datetime
from openpyxl import load_workbook
from pandas.api.types import CategoricalDtype

pd.options.display.float_format = "{:,.2f}".format


def generate_report(raw, format_file, dt):
    """FX5220 보고서 (2차) - RM 정보 추가 및 피벗 테이블 작성.

    통합 raw(.xlsx, 시트: cur_raw/bf_rm_info/cur_response) + 양식(format.xlsx)
    + 기준일(dt, yyyymmdd)을 받아 rm_info CSV를 저장하고, 양식 '입력표' 시트에
    피벗 결과를 기입하여 저장한다. 최종 파일명/경로는 하드코딩 규칙을 따른다.
    """
    logs = []

    def _log(m):
        logs.append(str(m))

    def _fail(step, e):
        import traceback
        _log(f"[ERROR] '{step}' 단계에서 오류: {e}")
        _log(traceback.format_exc())
        return {'ok': False, 'log': "\n".join(logs), 'outfile': None}

    try:
        _log("[1/5] 통합 raw 로드 중...")
        dt = str(dt)
        dt_trans = datetime.strptime(dt, '%Y%m%d')
        raw_file = pd.read_excel(raw, sheet_name=['cur_raw', 'bf_rm_info', 'cur_response'])
        output_dir = f"C:/Users/1598505/OneDrive - Standard Chartered Bank/1.보고서(WB)/2.BOK/2.월/FX5220_월보/2026년도/{dt[2:6]}/FX5220_rm_info_({dt[2:6]}기준).csv"
        outfile = f"C:/Users/1598505/OneDrive - Standard Chartered Bank/1.보고서(WB)/2.BOK/2.월/FX5220_월보/2026년도/{dt[2:6]}/FX5220_작성용({dt[2:6]}기준).xlsx"

        raw_copied = raw_file['cur_raw']
        bf_raw = raw_file['bf_rm_info']
        new_list_response = raw_file['cur_response']

        raw_copied.rename(columns={raw_copied.columns[0]: '계좌번호'}, inplace=True)
        bf_raw = bf_raw[['계좌번호', '업종', '기업규모', '용도', '담당자']].copy()
    except Exception as e:
        return _fail("통합 raw 로드", e)

    try:
        _log("[2/5] RM 정보 병합 및 rm_info 저장 중...")
        new_raw_response_copied = new_list_response.copy()
        for i in range(len(new_raw_response_copied)):
            if new_raw_response_copied.loc[i, '업종'] is None:
                new_raw_response_copied.loc[i, '업종'] = '비제조업'
                new_raw_response_copied.loc[i, '기업규모'] = '중소기업'
                new_raw_response_copied.loc[i, '용도'] = '해외사용운전자금'
                new_raw_response_copied.loc[i, '담당자'] = '오성진'

        raw_merged = raw_copied.merge(bf_raw, how='left', on='계좌번호')
        _log(f"  이번 달 raw data shape : {raw_merged.shape}")
        null_condition = raw_merged['담당자'].isnull()
        _log(f"  이번 달 raw data 중 신규 여신 shape : {new_raw_response_copied.shape}")
        _log(f"  이번 달 raw data 중 기존 여신 shape : {raw_merged[~null_condition].shape}")
        rm_info = pd.concat([raw_merged[~null_condition], new_raw_response_copied.drop(['base_dt'], axis=1)], axis=0)
        _log(f"  이번 달 raw data 신규 + 기존 여신 shape : {rm_info.shape}")
        if new_raw_response_copied.shape[0] + raw_merged[~null_condition].shape[0] != rm_info.shape[0]:
            _log('  [주의] 기존 여신과 신규 여신 합산 데이터의 행이 raw data의 행과 다릅니다')
        else:
            _log('  Pass')

        rm_info.to_csv(output_dir, index=False, encoding='utf-8-sig')
    except Exception as e:
        return _fail("RM 정보 병합/저장", e)

    try:
        _log("[3/5] 만기구간/통화분류 및 피벗 작성 중...")
        file_list = [raw_merged, new_raw_response_copied]

        for file in file_list:
            file['exec_amt'] = file['exec_amt'].astype('string').str.replace(',', '').astype('float')
            file['실행금액(US)'] = file['실행금액(US)'].astype('string').str.replace(',', '').astype('float')
            file['pybck_amt'] = file['pybck_amt'].astype('string').str.replace(',', '').astype('float')
            file['회수금액(US)'] = file['회수금액(US)'].astype('string').str.replace(',', '').astype('float')
            file['loan_asst_bs_amt'] = file['loan_asst_bs_amt'].astype('string').str.replace(',', '').astype('float')
            file['잔액(US)'] = file['잔액(US)'].astype('string').str.replace(',', '').astype('float')

            file['loan_deadln_dt'] = file['loan_deadln_dt'].astype('string').str.replace('-', '')
            file['loan_deadln_dt'] = pd.to_datetime(file['loan_deadln_dt'], format='%Y%m%d')

            file['new_start_dt'] = file['new_start_dt'].astype('string').str.replace('-', '')
            file['new_start_dt'] = pd.to_datetime(file['new_start_dt'], format='%Y%m%d')

            file['maturity1'] = file['loan_deadln_dt'] - file['new_start_dt']
            file['maturity2'] = file['loan_deadln_dt'] - dt_trans

            file['maturity1_gb'] = 0
            for i in range(len(file['maturity1'])):
                if file.loc[i, 'maturity1'].days < 91:
                    file.loc[i, 'maturity1_gb'] = 1
                elif file.loc[i, 'maturity1'].days < 186:
                    file.loc[i, 'maturity1_gb'] = 2
                elif file.loc[i, 'maturity1'].days < 371:
                    file.loc[i, 'maturity1_gb'] = 3
                elif file.loc[i, 'maturity1'].days < 741:
                    file.loc[i, 'maturity1_gb'] = 4
                else:
                    file.loc[i, 'maturity1_gb'] = 5

            file['maturity2_gb'] = 0
            for i in range(len(file['maturity2'])):
                if file.loc[i, 'maturity2'].days < 91:
                    file.loc[i, 'maturity2_gb'] = 1
                elif file.loc[i, 'maturity2'].days < 186:
                    file.loc[i, 'maturity2_gb'] = 2
                elif file.loc[i, 'maturity2'].days < 371:
                    file.loc[i, 'maturity2_gb'] = 3
                elif file.loc[i, 'maturity2'].days < 741:
                    file.loc[i, 'maturity2_gb'] = 4
                else:
                    file.loc[i, 'maturity2_gb'] = 5

            conditions = [(file['crncy_cd'].isin([1])), (file['crncy_cd'].isin([2])), (file['crncy_cd'].isin([161])), (~file['crncy_cd'].isin([1, 2, 161]))]
            choices = list(configs.CRNCY_CD_GROUP.keys())
            file['crncy_cd'] = np.select(conditions, choices, default='Unknown')

        new_raw_response_copied_v2 = new_raw_response_copied.drop(['base_dt'], axis=1)
        raw_merged_new = pd.concat([raw_merged[~null_condition], new_raw_response_copied_v2], axis=0)

        seg_order = ['대기업', '중소기업']
        indu_order = ['제조업', '비제조업']
        currency_order = ['미달러화', '일본엔화', '유로화', '기타통화']
        maturity_order = [1, 2, 3, 4, 5]
        funding_order = ['국내사용시설자금', '해외사용시설자금', '국내사용운전자금', '해외사용운전자금']

        raw_merged_new_v2 = raw_merged_new[['maturity1_gb', 'maturity2_gb', 'crncy_cd', '실행금액(US)', '기업규모', '업종', '용도', '회수금액(US)', '잔액(US)']]

        raw_merged_new_v2['기업규모'] = raw_merged_new_v2['기업규모'].astype(CategoricalDtype(categories=seg_order, ordered=True))
        raw_merged_new_v2['업종'] = raw_merged_new_v2['업종'].astype(CategoricalDtype(categories=indu_order, ordered=True))
        raw_merged_new_v2['crncy_cd'] = raw_merged_new_v2['crncy_cd'].astype(CategoricalDtype(categories=currency_order, ordered=True))
        raw_merged_new_v2['maturity1_gb'] = raw_merged_new_v2['maturity1_gb'].astype(CategoricalDtype(categories=maturity_order, ordered=True))
        raw_merged_new_v2['maturity2_gb'] = raw_merged_new_v2['maturity2_gb'].astype(CategoricalDtype(categories=maturity_order, ordered=True))
        raw_merged_new_v2['용도'] = raw_merged_new_v2['용도'].astype(CategoricalDtype(categories=funding_order, ordered=True))

        new_condition = raw_merged_new_v2['실행금액(US)'] > 0
        raw_new = raw_merged_new_v2[new_condition]

        raw_new_pivot_1st = pd.pivot_table(raw_new, values='실행금액(US)', index='crncy_cd', columns=['기업규모', '업종'], aggfunc='sum', observed=False, fill_value=0)
        raw_new_pivot_2nd = pd.pivot_table(raw_new, values='실행금액(US)', index='용도', columns=['기업규모', '업종'], aggfunc='sum', observed=False, fill_value=0)

        pybck_condition = raw_merged_new_v2['회수금액(US)'] > 0
        raw_pybck = raw_merged_new_v2[pybck_condition]

        raw_pybck_pivot_1st = pd.pivot_table(raw_pybck, values='회수금액(US)', index='crncy_cd', columns=['기업규모', '업종'], aggfunc='sum', observed=False, fill_value=0)
        raw_pybck_pivot_2nd = pd.pivot_table(raw_pybck, values='회수금액(US)', index='용도', columns=['기업규모', '업종'], aggfunc='sum', observed=False, fill_value=0)

        raw_merged_pivot_1st = pd.pivot_table(raw_merged_new_v2, values='잔액(US)', index='crncy_cd', columns=['기업규모', '업종'], aggfunc='sum', observed=False, fill_value=0)
        raw_merged_pivot_2nd = pd.pivot_table(raw_merged_new_v2, values='잔액(US)', index='용도', columns=['기업규모', '업종'], aggfunc='sum', observed=False, fill_value=0)

        raw_mature_pivot_1st = pd.pivot_table(raw_merged_new_v2, values='잔액(US)', index='maturity1_gb', columns=['기업규모', '업종'], aggfunc='sum', observed=False, fill_value=0)
        raw_mature_pivot_2nd = pd.pivot_table(raw_merged_new_v2, values='잔액(US)', index='maturity2_gb', columns=['기업규모', '업종'], aggfunc='sum', observed=False, fill_value=0)
    except Exception as e:
        return _fail("만기구간/통화분류/피벗", e)

    try:
        _log("[4/5] 양식(format) '입력표'에 데이터 기입 중...")
        wb = openpyxl.load_workbook(format_file, read_only=False)
        ws = wb['입력표']

        blocks = [
            (raw_new_pivot_1st, 5, 8, 2, 5),
            (raw_new_pivot_2nd, 15, 18, 2, 5),
            (raw_pybck_pivot_1st, 5, 8, 8, 11),
            (raw_pybck_pivot_2nd, 15, 18, 8, 11),
            (raw_merged_pivot_1st, 5, 8, 14, 17),
            (raw_merged_pivot_2nd, 15, 18, 14, 17),
            (raw_mature_pivot_1st, 5, 9, 20, 23),
            (raw_mature_pivot_2nd, 15, 19, 20, 23),
        ]
        for pivot_df, start_row, end_row, start_col, end_col in blocks:
            for r_idx, r in enumerate(range(start_row, end_row + 1)):
                for c_idx, c in enumerate(range(start_col, end_col + 1)):
                    if r_idx < pivot_df.shape[0] and c_idx < pivot_df.shape[1]:
                        ws.cell(row=r, column=c, value=pivot_df.iloc[r_idx, c_idx])
    except Exception as e:
        return _fail("양식 데이터 기입", e)

    try:
        _log("[5/5] 결과 저장 중...")
        wb.save(outfile)
        wb.close()
    except Exception as e:
        return _fail("결과 저장", e)

    _log(f"[완료] FX5220(2차) 보고서 작성 완료: {outfile}")
    return {'ok': True, 'log': "\n".join(logs), 'outfile': outfile, 'outfiles': [output_dir, outfile]}
