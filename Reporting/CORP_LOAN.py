import pandas as pd
import configs
import openpyxl
from datetime import datetime
from pandas.api.types import CategoricalDtype
import numpy as np


class corp_loan:
    def __init__(self):
        pass

    def _month_table(self, raw_data):
        print('--- 기업 여신 조사표 작성 시작 ---')
        print('--- raw data 편집 시작 ---')
        row_nm = ['F425', 'I551', 'I5613', 'R']
        num_list = []
        cols = raw_data.columns.tolist()
        for col in cols[1:]:
            raw_data[col] = raw_data[col] / 100000000
        print('--- 억단위 변환 완료 ---')
        for idx, nm in enumerate(row_nm):
            row_num = raw_data[raw_data['코드명'] == nm].index.to_numpy().item()
            num_list.append(row_num)
        print(f'--- 코드 {row_nm} 위치 파악 완료 : {num_list} ---')

        raw_parts = []
        for i in range(len(num_list)):
            if i == 0:
                raw_part = raw_data.loc[:num_list[i], :]
                raw_parts.append(raw_part)
            elif i == len(num_list) - 1:
                raw_part = raw_data.loc[num_list[i - 1] + 1:num_list[i], :]
                raw_parts.append(raw_part)
                raw_part = raw_data.loc[num_list[i] + 2:, :]
                raw_parts.append(raw_part)
            else:
                raw_part = raw_data.loc[num_list[i - 1] + 1:num_list[i], :]
                raw_parts.append(raw_part)

        additional_row_nm = ['F426', 'I55101', 'I55102', 'I55103', 'I55104', 'I55109', 'I5614']

        add_parts = {}

        for i, nm in enumerate(additional_row_nm):
            add_parts[nm] = [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]

        df_test = pd.DataFrame(add_parts)
        add_parts = df_test.T

        add_parts_new = add_parts.reset_index()
        add_parts_new.columns = raw_parts[0].columns

        raw_parts_integrated = pd.concat([raw_parts[0], add_parts_new.iloc[:1], raw_parts[1], add_parts_new.iloc[1:6], raw_parts[2], add_parts_new.iloc[6:], raw_parts[3], raw_parts[4]], ignore_index=True, axis=0)
        return raw_parts_integrated

    def complete_table(self, bf_raw_data, raw_data):
        # super().__init__() -> 부모 클래스의 속성을 물려받을 때에 사용
        bf_table = self._month_table(bf_raw_data)
        table = self._month_table(raw_data)

        bf_cols = ['대기업 전월말잔액', '중소기업 전월말잔액', '개인사업자 전월말잔액']
        cols = ['대기업 금월말잔액', '중소기업 금월말잔액', '개인사업자 금월말잔액']

        for i, bf_col in enumerate(bf_cols):
            table[bf_col] = bf_table[cols[i]]
        table.set_index('코드명', inplace=True)
        return table

    def fs_00409(self, raw_data):
        raw_data['등급'] = raw_data['등급'].fillna('무등급')

        conditions = [(raw_data['등급'].isin(['1A', '1B', '2A', '2B'])), (raw_data['등급'].isin(['3A', '3B', '4A'])),
                      (raw_data['등급'].isin(['4B', '5A', '5B'])), (raw_data['등급'].isin(['6A', '6B', '7A', '7B', '8A', '8B'])), (raw_data['등급'].isin(['9A', '9B', '10A', '10B'])),
                      (raw_data['등급'].isin(['11A', '11B', '11C'])), (raw_data['등급'].isin(['12A', '12B', '12C'])), (raw_data['등급'].isin(['13'])),
                      (raw_data['등급'].isin(['14A'])), (raw_data['등급'].isin(['14B'])), (raw_data['등급'].isin(['무등급'])), (raw_data['등급'].isin(['소매 익스포저']))]
        choices = list([1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12])
        raw_data['new등급'] = np.select(conditions, choices, default=8)

        new_grade_order = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12]
        raw_data['new등급'] = raw_data['new등급'].astype(CategoricalDtype(categories=new_grade_order, ordered=True))

        seg_order = ['대기업', '중소+개인', '개인']
        raw_data['분류'] = raw_data['분류'].astype(CategoricalDtype(categories=seg_order, ordered=True))

        raw_data['KGAAP 신규금액'] = raw_data['KGAAP 신규금액'] / 100000000
        raw_data['월말 금액'] = raw_data['월말 금액'] / 100000000

        raw_pivot = pd.pivot_table(
            raw_data,
            values=['KGAAP 신규금액', '월말 금액'],
            index='new등급',
            columns=['분류'],
            aggfunc='sum',
            fill_value=0)

        return raw_pivot

    def fs_00410(self, raw_data):
        raw = raw_data[['기업규모', '만기도래', '만기상환', '중도상환', '현재잔액']]
        raw.set_index('기업규모', inplace=True)

        raw = raw / 100000000
        raw.loc['기업'] = 0

        raw.iloc[1, :] = raw.iloc[1, :] + raw.iloc[2, :]
        raw.iloc[3, :] = raw.iloc[0, :] + raw.iloc[1, :]
        raw['만기연장액'] = raw['만기도래'] - raw['만기상환'] - raw['중도상환'] - raw['현재잔액']
        raw['만기연장률'] = raw['만기연장액'] / (raw['만기도래'] - raw['중도상환'] - raw['현재잔액'])
        raw['만기연장률'] = raw['만기연장률'] * 100
        raw.index = raw.index.astype(pd.CategoricalDtype(categories=['기업', '대기업', '중소기업', '개인사업자'], ordered=True))
        raw = raw.sort_index()

        return raw


def processing(ws, start_row, end_row, source_column, target_column):
    for row in range(start_row, end_row):
        source_val = ws.cell(row=row, column=source_column).value
        ws.cell(row=row, column=target_column).value = source_val
    return ws


def input_processing(df, ws, start_row, end_row, start_col, end_col):
    for r_idx, r in enumerate(range(start_row, end_row + 1)):
        for c_idx, c in enumerate(range(start_col, end_col + 1)):
            if r_idx < df.shape[0] and c_idx < df.shape[1]:
                ws.cell(row=r, column=c, value=df.iloc[r_idx, c_idx])
    return ws


def processing_00409(df, start_row, end_row, source_column, ws, start_t_row, end_t_row, target_column):
    for row, t_row in zip(range(start_row, end_row + 1), range(start_t_row, end_t_row + 1)):
        source_val = df.iloc[row, source_column]
        ws.cell(row=t_row, column=target_column).value = source_val
    return ws


def generate_report(raw, format_file, dt):
    """기업 여신 조사표 (FS_00401/402/409/410).

    통합 raw(.xlsx, 다중 시트) + 양식(format.xlsx) + 기준년월(dt, YYYYMM)을 받아
    양식의 FS00401/FS00402/FS00409/FS00410 시트에 최신 데이터를 기입하여 저장한다.
    최종 파일명/경로는 기존 하드코딩 규칙을 그대로 따른다.
    """
    logs = []

    def _log(m):
        logs.append(str(m))

    def _fail(step, e):
        import traceback
        _log(f"[ERROR] '{step}' 단계에서 오류: {e}")
        _log(traceback.format_exc())
        return {'ok': False, 'log': "\n".join(logs), 'outfile': None}

    try:
        _log("[1/4] 통합 raw 로드 중...")
        dt = str(dt)
        dt_trans = datetime.strptime(dt, '%Y%m')
        dt_bf = f"{dt_trans.year}{dt_trans.month - 1:02d}"
        _log(f"  이전 기준년월: {dt_bf}")
        raw_file = pd.read_excel(raw, sheet_name=[f'fs_00401_{dt_bf}', f'fs_00401_{dt}', f'fs_00409_{dt}', f'fs_00410_{dt}', '원화대출금_FS00402'])
        outfile = f'C:/Users/1598505/OneDrive - Standard Chartered Bank/1.보고서(WB)/2.BOK/2.월/기업여신조사표/2026년도/{dt[2:]}/기업여신조사표_{dt[2:]}_SC제일은행.xlsx'

        raw_copied = raw_file[f'fs_00401_{dt}']
        bf_raw_copied = raw_file[f'fs_00401_{dt_bf}']
        raw_copied2 = raw_file[f'fs_00409_{dt}']
        raw_copied3 = raw_file[f'fs_00410_{dt}']
    except Exception as e:
        return _fail("통합 raw 로드", e)

    try:
        _log("[2/4] FS 테이블 산출 중 (00401/00402/00409/00410)...")
        corp = corp_loan()
        fs_00401 = corp.complete_table(bf_raw_copied, raw_copied)
        fs_00402 = raw_file['원화대출금_FS00402'].iloc[:, 2:]
        fs_00409 = corp.fs_00409(raw_copied2)
        fs_00410 = corp.fs_00410(raw_copied3)
    except Exception as e:
        return _fail("FS 테이블 산출", e)

    try:
        _log("[3/4] 양식(format)에 데이터 기입 중...")
        wb = openpyxl.load_workbook(format_file, read_only=False)
        ws1 = wb['FS00401']
        ws1 = input_processing(fs_00401, ws1, 8, 146, 3, 21)
        ws2 = wb['FS00402']
        ws2 = input_processing(fs_00402, ws2, 8, 146, 3, 23)

        ws3 = wb['FS00409']
        ws3 = processing(ws3, 7, 19, 8, 6)
        ws3 = processing(ws3, 7, 19, 13, 9)
        ws3 = processing(ws3, 7, 19, 14, 10)
        ws3 = processing_00409(df=fs_00409, start_row=0, end_row=11, source_column=0, ws=ws3, start_t_row=7, end_t_row=19, target_column=7)
        ws3 = processing_00409(df=fs_00409, start_row=0, end_row=11, source_column=3, ws=ws3, start_t_row=7, end_t_row=19, target_column=8)
        ws3 = processing_00409(df=fs_00409, start_row=0, end_row=11, source_column=1, ws=ws3, start_t_row=7, end_t_row=19, target_column=11)
        ws3 = processing_00409(df=fs_00409, start_row=0, end_row=11, source_column=2, ws=ws3, start_t_row=7, end_t_row=19, target_column=12)
        ws3 = processing_00409(df=fs_00409, start_row=0, end_row=11, source_column=4, ws=ws3, start_t_row=7, end_t_row=19, target_column=13)
        ws3 = processing_00409(df=fs_00409, start_row=0, end_row=11, source_column=5, ws=ws3, start_t_row=7, end_t_row=19, target_column=14)

        ws4 = wb['FS00410']
        ws4 = input_processing(fs_00410, ws4, 5, 8, 3, 8)
    except Exception as e:
        return _fail("양식 데이터 기입", e)

    try:
        _log("[4/4] 결과 저장 중...")
        wb.save(outfile)
        wb.close()
    except Exception as e:
        return _fail("결과 저장", e)

    _log(f"[완료] 기업 여신 조사표 작성 완료: {outfile}")
    return {'ok': True, 'log': "\n".join(logs), 'outfile': outfile}
