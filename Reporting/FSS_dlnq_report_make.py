import pandas as pd
import openpyxl
from pathlib import Path


def processing(df, start_row, end_row, source_column, ws, start_t_row, end_t_row, target_column):
    for row, t_row in zip(range(start_row, end_row + 1), range(start_t_row, end_t_row + 1)):
        source_val = df.iloc[row, source_column]
        ws.cell(row=t_row, column=target_column).value = source_val
    return ws


def generate_report(raw, format_file, yymmdd):
    """FSS 연체 보고서.

    통합 raw(.xlsx, 시트 2종) + 양식(format.xlsx)을 받아 '(붙임1) 요약'·
    '(붙임2) 업종별 현황' 시트에 최신 데이터를 기입하여 저장한다.
    최종 파일명/경로는 기존 하드코딩 규칙을 그대로 따른다.
    """
    logs = []

    def _log(m):
        logs.append(str(m))

    def _fail(step, e):
        import traceback
        _log(f"[ERROR] '{step}' 단계에서 오류: {e}")
        _log(traceback.format_exc())
        return {'ok': False, 'log': "\n".join(logs), 'outfile': None}

    try:
        _log("[1/4] 통합 raw 로드 중...")
        raw_file = pd.read_excel(raw, sheet_name=['lmla_fss_corp_dlnq_sts_dtl_1', 'lmla_fss_corp_dlnq_sts_dtl_2'])
        raw1_copied = raw_file['lmla_fss_corp_dlnq_sts_dtl_1']
        raw1_final = raw1_copied.drop(index=[0, 1])

        raw2_copied = raw_file['lmla_fss_corp_dlnq_sts_dtl_2']
        col_list = raw2_copied.columns.tolist()[3:]
    except Exception as e:
        return _fail("통합 raw 로드", e)

    try:
        _log("[2/4] 연체 데이터 보정 중...")
        for c in col_list:
            raw2_copied.loc[17, c] += (raw2_copied.loc[14, c] - raw2_copied.loc[15, c])
            raw2_copied.loc[36, c] += (raw2_copied.loc[33, c] - raw2_copied.loc[34, c])
            raw2_copied.loc[55, c] += (raw2_copied.loc[52, c] - raw2_copied.loc[53, c])
            raw2_copied.loc[74, c] += (raw2_copied.loc[71, c] - raw2_copied.loc[72, c])

        raw2_final = raw2_copied.drop(index=[14, 33, 52, 71])
    except Exception as e:
        return _fail("연체 데이터 보정", e)

    try:
        _log("[3/4] 양식(format)에 데이터 기입 중...")
        wb = openpyxl.load_workbook(format_file, read_only=False)
        ws1 = wb['(붙임 1) 요약(대출잔액, 연체잔액, 신규연체)']
        ws2 = wb['(붙임 2) 기업대출 업종별 현황']

        r = 11
        last_col = 0
        for col in range(ws1.max_column, 0, -1):
            if ws1.cell(row=r, column=col).value is not None:
                last_col = col + 1
                break

        last_col2 = 0
        for col in range(ws2.max_column, 0, -1):
            if ws2.cell(row=r, column=col).value is not None:
                last_col2 = col + 1
                break

        processing(df=raw1_final, start_row=0, end_row=6, source_column=3, ws=ws1, start_t_row=11, end_t_row=17, target_column=last_col)
        processing(df=raw1_final, start_row=0, end_row=6, source_column=4, ws=ws1, start_t_row=30, end_t_row=36, target_column=last_col)
        processing(df=raw1_final, start_row=0, end_row=6, source_column=5, ws=ws1, start_t_row=49, end_t_row=55, target_column=last_col)

        processing(df=raw2_final, start_row=0, end_row=17, source_column=3, ws=ws2, start_t_row=11, end_t_row=28, target_column=last_col2)
        processing(df=raw2_final, start_row=0, end_row=17, source_column=4, ws=ws2, start_t_row=106, end_t_row=123, target_column=last_col2)
        processing(df=raw2_final, start_row=0, end_row=17, source_column=5, ws=ws2, start_t_row=201, end_t_row=218, target_column=last_col2)

        processing(df=raw2_final, start_row=18, end_row=35, source_column=3, ws=ws2, start_t_row=34, end_t_row=51, target_column=last_col2)
        processing(df=raw2_final, start_row=18, end_row=35, source_column=4, ws=ws2, start_t_row=129, end_t_row=146, target_column=last_col2)
        processing(df=raw2_final, start_row=18, end_row=35, source_column=5, ws=ws2, start_t_row=224, end_t_row=241, target_column=last_col2)

        processing(df=raw2_final, start_row=36, end_row=53, source_column=3, ws=ws2, start_t_row=57, end_t_row=74, target_column=last_col2)
        processing(df=raw2_final, start_row=36, end_row=53, source_column=4, ws=ws2, start_t_row=152, end_t_row=169, target_column=last_col2)
        processing(df=raw2_final, start_row=36, end_row=53, source_column=5, ws=ws2, start_t_row=247, end_t_row=264, target_column=last_col2)

        processing(df=raw2_final, start_row=54, end_row=71, source_column=3, ws=ws2, start_t_row=80, end_t_row=97, target_column=last_col2)
        processing(df=raw2_final, start_row=54, end_row=71, source_column=4, ws=ws2, start_t_row=175, end_t_row=192, target_column=last_col2)
        processing(df=raw2_final, start_row=54, end_row=71, source_column=5, ws=ws2, start_t_row=270, end_t_row=287, target_column=last_col2)
    except Exception as e:
        return _fail("양식 데이터 기입", e)

    try:
        _log("[4/4] 결과 저장 중...")
        outfile = f'C:/Users/1598505/OneDrive - Standard Chartered Bank/1.보고서(WB)/1.FSS/1.업무보고서/1.주보_10일보/2.10일보/{yymmdd}/(260313 통합양식)_국내은행 연체율({yymmdd}기준)_SC제일은행.xlsx'
        wb.save(outfile)
        wb.close()
    except Exception as e:
        return _fail("결과 저장", e)

    _log(f"[완료] FSS 연체 보고서 작성 완료: {outfile}")
    return {'ok': True, 'log': "\n".join(logs), 'outfile': outfile}
